"""Read an ECI Form 20 that has already been extracted into spreadsheet tables.

The Giridih Form 20s for VS-2019 and VS-2024 (db/seed/form20/) arrive as one
worksheet per printed page, `Page_1 ... Page_12`. Each sheet repeats a header:

    row 1  column numbers ("0.0", "1.0", ...) - an extraction artefact
    row 2  "Serial No. Of Polling Station", "No of Valid Votes Cast in favour of",
           "Total of Valid Votes", "No. Of Rejected Votes", "NOTA", "Total",
           "No. Of Tendered Votes"
    row 3  the candidate names, one column each

followed by one row per polling station:

    serial, ps_number, <votes per candidate...>, valid, rejected, NOTA, total, tendered

where **valid excludes NOTA** (the Form 20 convention) and total = valid + NOTA
+ rejected. The last page ends with three footer rows - "Total EVM Votes",
"Total Postal Ballot Votes" and "Total Votes Polled" - with the same columns.

Two extraction defects are handled here rather than by hand-editing the file:

* On the 2019 pages the last station of every page is interleaved with the
  "Page" stamp printed at the foot of the sheet: the serial reads "P3a2", the PS
  number "g3e2" and the first candidate cell "1 2" (page number 1, value 2) for
  station 32. The station number is the digits of the first two cells (they must
  agree) and the value is the last token. Row arithmetic then has to hold, so a
  wrong recovery cannot pass silently.
* The 2024 file repeats "Total Votes Polled" on its own final page with the
  columns shifted one to the left. Footer values are therefore read as the
  numeric cells after the label, wherever they start.

Nothing here touches the database: `read()` parses, `problems()` lists every
integrity failure, and the loader (ingest/load_form20_tables.py) refuses to
write anything while that list is non-empty.
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path

TAIL = ("valid", "rejected", "nota", "total", "tendered")
FOOTERS = {
    "total evm votes": "evm",
    "total postal ballot votes": "postal",
    "total votes polled": "polled",
}
_INT = re.compile(r"-?\d+(?:\.0+)?")


@dataclass
class Totals:
    """One footer row: votes per candidate plus the five tail columns."""

    votes: list[int]
    valid: int          # excludes NOTA, as printed
    rejected: int
    nota: int
    total: int
    tendered: int


@dataclass
class Booth:
    page: int
    serial: int
    ps_number: int
    votes: list[int]
    valid: int          # excludes NOTA, as printed
    rejected: int
    nota: int
    total: int
    tendered: int
    recovered: bool = False   # rebuilt from a row garbled by the page stamp

    @property
    def valid_incl_nota(self) -> int:
        """The project's definition of valid votes (METRICS.md): NOTA included."""
        return self.valid + self.nota


@dataclass
class Form20Table:
    path: Path
    sha256: str
    candidates: list[str]
    booths: list[Booth]
    footers: dict[str, Totals] = field(default_factory=dict)
    pages: int = 0

    @property
    def evm(self) -> Totals | None:
        return self.footers.get("evm")

    @property
    def postal(self) -> Totals | None:
        return self.footers.get("postal")

    @property
    def polled(self) -> Totals | None:
        return self.footers.get("polled")

    def published_votes(self) -> dict[str, int]:
        """Total votes per candidate including postal ballots - the declared result."""
        if self.polled is None:
            raise ValueError(f"{self.path.name}: no 'Total Votes Polled' row")
        return dict(zip(self.candidates, self.polled.votes, strict=True))

    def booth_by_ps(self) -> dict[int, Booth]:
        return {b.ps_number: b for b in self.booths}


def clean_name(raw: str) -> str:
    """'NIRBHAY\\nKUMAR\\nSHAHABADI' -> 'NIRBHAY KUMAR SHAHABADI'."""
    return re.sub(r"\s+", " ", str(raw or "")).strip()


def name_key(name: str) -> str:
    """Case- and punctuation-insensitive key for matching a header to a seeded candidate."""
    return re.sub(r"[^a-z]", "", clean_name(name).lower())


def display_name(name: str) -> str:
    """Title case for display ('DR. BARNABAS HEMBROM' -> 'Dr. Barnabas Hembrom')."""
    return " ".join(part.capitalize() for part in clean_name(name).split(" "))


def _int(value) -> int:
    text = str(value).strip()
    if not _INT.fullmatch(text):
        raise ValueError(f"not an integer: {value!r}")
    return int(float(text))


def _digits(value) -> int:
    found = re.sub(r"\D", "", str(value or ""))
    if not found:
        raise ValueError(f"no station number in {value!r}")
    return int(found)


def _sheet_rows(path: Path) -> list[tuple[str, list[list]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)])
                for ws in workbook.worksheets]
    finally:
        workbook.close()


def _header_index(rows: list[list]) -> int:
    """Index of the row holding the candidate names (the one after 'Serial ...')."""
    for i, row in enumerate(rows):
        if row and row[0] and str(row[0]).strip().lower().startswith("serial"):
            return i + 1
    raise ValueError("no 'Serial No. Of Polling Station' header row")


def read(path: str | Path) -> Form20Table:
    """Parse every sheet of an extracted Form 20 workbook."""
    path = Path(path)
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    sheets = _sheet_rows(path)

    candidates: list[str] | None = None
    booths: list[Booth] = []
    footers: dict[str, Totals] = {}

    for page_no, (_title, rows) in enumerate(sheets, 1):
        try:
            h = _header_index(rows)
        except ValueError:
            h = None                      # a page holding only footer rows
        if h is not None:
            # Names are the non-empty cells after the serial column: columns 2.. on
            # most pages, 1.. on the 2024 file's final page, which is shifted left.
            names = [clean_name(c) for c in rows[h][1:] if c not in (None, "")]
            if candidates is None:
                candidates = names
            elif names != candidates:
                raise ValueError(f"{path.name} page {page_no}: candidate header differs "
                                 f"from page 1 ({names} vs {candidates})")
            body = rows[h + 1:]
        else:
            body = rows
        if candidates is None:
            raise ValueError(f"{path.name}: no candidate header on page {page_no}")
        n = len(candidates)

        for row in body:
            if not row or all(c in (None, "") for c in row):
                continue
            label = clean_name(row[0]).lower()
            if label.startswith("total"):
                key = FOOTERS.get(label)
                if key is None:
                    raise ValueError(f"{path.name} page {page_no}: unknown footer {row[0]!r}")
                values = [_int(c) for c in row[1:] if c not in (None, "")]
                if len(values) != n + len(TAIL):
                    raise ValueError(f"{path.name}: footer {row[0]!r} has {len(values)} values, "
                                     f"expected {n + len(TAIL)}")
                footers[key] = Totals(values[:n], *values[n:])
                continue
            if label.replace(".0", "").isdigit() or _INT.fullmatch(str(row[0]).strip()):
                serial, ps, first, recovered = _int(row[0]), _int(row[1]), row[2], False
            else:
                serial, ps = _digits(row[0]), _digits(row[1])
                if serial != ps:
                    raise ValueError(f"{path.name} page {page_no}: garbled row {row[:3]} - "
                                     f"serial {serial} and PS {ps} disagree")
                first, recovered = str(row[2]).split()[-1], True
            cells = [first, *row[3:2 + n]]
            tail = row[2 + n:2 + n + len(TAIL)]
            booths.append(Booth(page_no, serial, ps, [_int(c) for c in cells],
                                *[_int(c) for c in tail], recovered=recovered))

    return Form20Table(path, sha, candidates or [], booths, footers, len(sheets))


def problems(table: Form20Table) -> list[str]:
    """Every integrity failure. Empty means the table can be loaded."""
    out: list[str] = []
    n = len(table.candidates)
    if n < 2:
        out.append(f"only {n} candidate column(s) found")

    numbers = [b.ps_number for b in table.booths]
    if len(numbers) != len(set(numbers)):
        dupes = sorted({x for x in numbers if numbers.count(x) > 1})
        out.append(f"duplicate polling stations: {dupes[:10]}")
    if numbers and sorted(set(numbers)) != list(range(1, max(numbers) + 1)):
        missing = sorted(set(range(1, max(numbers) + 1)) - set(numbers))
        out.append(f"polling stations missing from the sequence: {missing[:10]}")

    for b in table.booths:
        if len(b.votes) != n:
            out.append(f"PS {b.ps_number}: {len(b.votes)} vote cells for {n} candidates")
            continue
        if any(v < 0 for v in b.votes):
            out.append(f"PS {b.ps_number}: negative votes")
        if sum(b.votes) != b.valid:
            out.append(f"PS {b.ps_number}: candidates sum to {sum(b.votes)}, "
                       f"printed valid is {b.valid}")
        if b.valid + b.nota + b.rejected != b.total:
            out.append(f"PS {b.ps_number}: valid {b.valid} + NOTA {b.nota} + rejected "
                       f"{b.rejected} != total {b.total}")

    for key in ("evm", "postal", "polled"):
        if key not in table.footers:
            out.append(f"footer row missing: {key}")
    if table.evm is not None:
        columns = [sum(b.votes[i] for b in table.booths) for i in range(n)]
        if columns != table.evm.votes:
            bad = [table.candidates[i] for i in range(n) if columns[i] != table.evm.votes[i]]
            out.append(f"booth sums differ from 'Total EVM Votes' for {bad}")
        for attr in TAIL:
            got = sum(getattr(b, attr) for b in table.booths)
            if got != getattr(table.evm, attr):
                out.append(f"booth {attr} sum {got} != 'Total EVM Votes' {getattr(table.evm, attr)}")
    if table.evm and table.postal and table.polled:
        for i, name in enumerate(table.candidates):
            if table.evm.votes[i] + table.postal.votes[i] != table.polled.votes[i]:
                out.append(f"{name}: EVM + postal != total polled")
        for attr in TAIL:
            if getattr(table.evm, attr) + getattr(table.postal, attr) != getattr(table.polled, attr):
                out.append(f"{attr}: EVM + postal != total polled")
    return out
