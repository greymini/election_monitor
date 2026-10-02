"""The real Giridih Form 20 (VS-2019, VS-2024) loaded through ingest.load_form20_tables.

A separate database on the same server: the shared e2e dataset is synthetic
and occupies AC 32, and this one must start with no booths so the loader's
--create-booths path is what runs. Every expected figure below is the declared
result printed in the Form 20 itself (db/seed/form20/).
"""

from __future__ import annotations

import os
import shutil

import pytest

from tests.e2e.conftest import requires_db
from tests.paths import BACKEND

pytestmark = requires_db

FORM20 = BACKEND / "db" / "seed" / "form20"
F2024 = FORM20 / "giridih_vs2024_form20.xlsx"
F2019 = FORM20 / "giridih_vs2019_form20.xlsx"
DB_NAME = "giridih_form20_test"


def _with_db(url: str, name: str) -> str:
    from psycopg.conninfo import conninfo_to_dict, make_conninfo

    params = conninfo_to_dict(url)
    params["dbname"] = name
    return make_conninfo(**params)


def _use(url: str) -> None:
    from common.config import get_settings
    from common.db import close_pools

    os.environ["DATABASE_URL"] = url
    get_settings.cache_clear()
    close_pools()


@pytest.fixture(scope="module")
def real(db_url):
    import psycopg

    from analytics.refresh import VIEWS
    from db import apply_migrations
    from ingest.load_form20_tables import load_table

    previous = os.environ.get("DATABASE_URL")
    with psycopg.connect(db_url, autocommit=True) as admin:
        admin.execute(f'DROP DATABASE IF EXISTS "{DB_NAME}"')
        admin.execute(f'CREATE DATABASE "{DB_NAME}"')
    url = _with_db(db_url, DB_NAME)
    _use(url)
    try:
        assert apply_migrations.main(["--seed"]) == 0
        summaries = {
            "VS-2024": load_table(F2024, 32, "VS-2024", create_booths=True),
            "VS-2019": load_table(F2019, 32, "VS-2019"),
        }
        conn = psycopg.connect(url, autocommit=True, row_factory=psycopg.rows.dict_row)
        for view in VIEWS:
            conn.execute(f"REFRESH MATERIALIZED VIEW {view}")
        yield {"url": url, "conn": conn, "summaries": summaries}
        conn.close()
    finally:
        _use(previous or db_url)


def one(conn, sql, params=()):
    return conn.execute(sql, params).fetchone()


def test_every_station_has_a_booth_and_two_elections(real):
    c = real["conn"]
    assert one(c, "SELECT COUNT(*) AS n FROM booth b JOIN ac a USING (ac_id) "
                  "WHERE a.ac_number = 32")["n"] == 367
    rows = c.execute("SELECT election_label, COUNT(*) AS n FROM mv_result_booth_wide "
                     "GROUP BY 1 ORDER BY 1").fetchall()
    assert {r["election_label"]: r["n"] for r in rows} == {"VS-2019": 367, "VS-2024": 367}


@pytest.mark.parametrize("label, expected", [
    ("VS-2024", dict(winner_candidate="Sudivya Kumar", winner_party="JMM", winner_votes=94042,
                     runner_candidate="Nirbhay Kumar Shahabadi", runner_party="BJP",
                     runner_votes=90204, margin_votes=3838, valid_votes=207682,
                     votes_polled=207821, nota=2004, rejected=139, postal_votes=1905,
                     evm_votes=205777, contestants=14, electors=304898,
                     electors_source="published")),
    ("VS-2019", dict(winner_candidate="Sudivya Kumar", winner_party="JMM", winner_votes=80871,
                     runner_candidate="Nirbhay Kumar Shahabadi", runner_party="BJP",
                     runner_votes=64987, margin_votes=15884, valid_votes=167822,
                     votes_polled=167893, nota=2773, rejected=71, postal_votes=271,
                     evm_votes=167551, contestants=12, electors=264814,
                     electors_source="published")),
])
def test_the_constituency_result_is_the_declared_one(real, label, expected):
    row = one(real["conn"], "SELECT * FROM mv_ac_summary WHERE election_label = %s", (label,))
    got = {k: row[k] for k in expected}
    assert got == expected


def test_margin_percent_and_turnout(real):
    row = one(real["conn"], "SELECT margin_pct, turnout_pct FROM mv_ac_summary "
                            "WHERE election_label = 'VS-2024'")
    assert float(row["margin_pct"]) == 1.85          # 3,838 / 207,682
    assert float(row["turnout_pct"]) == 68.16        # 207,821 / 304,898 (published electors)


@pytest.mark.parametrize("label, led", [
    ("VS-2024", {"BJP": 193, "JMM": 166, "JLKM": 8}),
    ("VS-2019", {"JMM": 189, "BJP": 178}),
])
def test_booths_led(real, label, led):
    rows = real["conn"].execute(
        "SELECT winner_party, COUNT(*) AS n FROM mv_result_booth_wide "
        "WHERE election_label = %s GROUP BY 1", (label,)).fetchall()
    assert {r["winner_party"]: r["n"] for r in rows} == led


def test_candidates_without_a_recorded_party_are_kept_apart(real):
    rows = real["conn"].execute(
        "SELECT DISTINCT contestant FROM mv_result_booth_candidate "
        "WHERE election_label = 'VS-2024' AND party = 'UNK'").fetchall()
    assert len(rows) == 11
    assert all(r["contestant"].startswith("UNK:") for r in rows)


def test_swing_matches_a_hand_computation(real):
    """PS 1: JMM share of valid votes (NOTA included), 2024 minus 2019."""
    from ingest import form20_tables

    t24, t19 = form20_tables.read(F2024), form20_tables.read(F2019)
    b24, b19 = t24.booth_by_ps()[1], t19.booth_by_ps()[1]
    i24, i19 = t24.candidates.index("SUDIVYA KUMAR"), t19.candidates.index("SUDIVYA KUMAR")
    share24 = 100 * b24.votes[i24] / b24.valid_incl_nota
    share19 = 100 * b19.votes[i19] / b19.valid_incl_nota
    row = one(real["conn"],
              "SELECT s.swing_pct FROM mv_swing s JOIN booth b ON b.booth_uid = s.booth_uid "
              "WHERE s.election_label = 'VS-2024' AND s.party = 'JMM' AND b.current_ps_number = 1")
    assert row is not None, "swing withheld - the 2019 crosswalk is not trusted"
    assert float(row["swing_pct"]) == pytest.approx(share24 - share19, abs=0.01)
    n = one(real["conn"], "SELECT COUNT(*) AS n FROM mv_swing WHERE election_label = 'VS-2024' "
                          "AND party = 'JMM' AND swing_pct IS NOT NULL")["n"]
    assert n == 367


def test_the_2019_mapping_records_its_evidence(real):
    evidence = real["summaries"]["VS-2019"]["numbering_evidence"]
    assert evidence["stable"] and evidence["same_ps_r"] >= 0.9
    row = one(real["conn"], "SELECT COUNT(*) AS n, MIN(note) AS note FROM crosswalk_audit")
    assert row["n"] == 367 and "r =" in row["note"]


def test_the_sources_are_registered_as_real(real):
    rows = real["conn"].execute(
        "SELECT filename, is_synthetic, parse_status FROM source_doc WHERE kind = 'form20' "
        "ORDER BY filename").fetchall()
    assert [(r["filename"], r["is_synthetic"], r["parse_status"]) for r in rows] == [
        ("giridih_vs2019_form20.xlsx", False, "loaded"),
        ("giridih_vs2024_form20.xlsx", False, "loaded")]
    page = one(real["conn"], "SELECT source_doc, source_page FROM result_booth_meta m "
                             "JOIN election e USING (election_id) "
                             "WHERE e.label = 'VS-2024' AND m.ps_number = 40")
    assert (page["source_doc"], page["source_page"]) == ("giridih_vs2024_form20.xlsx", 2)


def test_the_winner_is_flagged(real):
    rows = real["conn"].execute(
        "SELECT e.label, c.name_en FROM candidate c JOIN election e USING (election_id) "
        "WHERE c.is_winner ORDER BY 1").fetchall()
    assert [(r["label"], r["name_en"]) for r in rows] == [
        ("VS-2019", "Sudivya Kumar"), ("VS-2024", "Sudivya Kumar")]


def test_validation_passes_on_the_real_load(real):
    from ingest import validate

    totals = validate.check_form20_totals()[0]
    rows = validate.check_row_arithmetic()[0]
    assert totals.passed, totals.rows
    assert rows.passed, rows.rows


def test_reloading_is_idempotent(real):
    from ingest.load_form20_tables import load_table

    c = real["conn"]
    before = one(c, "SELECT (SELECT COUNT(*) FROM candidate) AS cands, "
                    "(SELECT COUNT(*) FROM result_booth) AS rows, "
                    "(SELECT SUM(votes) FROM result_booth) AS votes")
    load_table(F2024, 32, "VS-2024")
    after = one(c, "SELECT (SELECT COUNT(*) FROM candidate) AS cands, "
                   "(SELECT COUNT(*) FROM result_booth) AS rows, "
                   "(SELECT SUM(votes) FROM result_booth) AS votes")
    assert after == before


def test_a_tampered_file_is_refused_and_nothing_changes(real, tmp_path):
    from openpyxl import load_workbook

    from ingest.load_form20_tables import LoadRefused, load_table

    bad = tmp_path / "tampered.xlsx"
    shutil.copy(F2024, bad)
    wb = load_workbook(bad)
    wb["Page_1"].cell(row=4, column=4).value = "448"
    wb.save(bad)
    c = real["conn"]
    before = one(c, "SELECT SUM(votes) AS v FROM result_booth")["v"]
    with pytest.raises(LoadRefused):
        load_table(bad, 32, "VS-2024")
    assert one(c, "SELECT SUM(votes) AS v FROM result_booth")["v"] == before


# ---------------------------------------------------------------------------
# The API on the real load: what the pages read.
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def api(real):
    from fastapi.testclient import TestClient

    from api.deps import hash_password
    from api.main import app
    from common.config import get_settings
    from common.db import query_one

    os.environ.setdefault("JWT_SECRET", "contract-test-secret-at-least-32-bytes-long")
    get_settings.cache_clear()
    block = query_one("SELECT block_id FROM block WHERE name_en = 'PS list not loaded'")
    headers = {}
    client = TestClient(app, raise_server_exceptions=False)
    for i, (role, block_id) in enumerate((("strategist", None),
                                          ("block", block["block_id"]))):
        phone = f"96000000{i:02d}"
        query_one("INSERT INTO app_user (phone, name, role, block_id, password_hash, "
                  "daily_token_budget) VALUES (%s, %s, %s, %s, %s, 1000) "
                  "ON CONFLICT (phone) DO UPDATE SET role = EXCLUDED.role "
                  "RETURNING user_id",
                  (phone, f"Form 20 {role}", role, block_id, hash_password("pw-" + role)))
        token = client.post("/auth/login", json={"phone": phone, "password": "pw-" + role})
        assert token.status_code == 200, token.text
        headers[role] = {"Authorization": f"Bearer {token.json()['access_token']}"}
    return client, headers


def _get(api, path, role="strategist"):
    client, headers = api
    response = client.get(path, headers=headers[role])
    assert response.status_code == 200, response.text[:500]
    return response.json()


def test_summary_reports_the_declared_result_with_postal_ballots(api):
    body = _get(api, "/acs/32/summary")
    assert body["synthetic"] is False
    rows = {e["label"]: e for e in body["elections"]}
    e24 = rows["VS-2024"]
    assert (e24["winner_candidate"], e24["winner_votes"]) == ("Sudivya Kumar", 94042)
    assert (e24["runner_candidate"], e24["runner_votes"]) == ("Nirbhay Kumar Shahabadi", 90204)
    assert (e24["margin_votes"], float(e24["margin_pct"])) == (3838, 1.85)
    assert (e24["evm_votes"], e24["postal_votes"], e24["rejected"]) == (205777, 1905, 139)
    assert (e24["total_valid"], e24["votes"], e24["contestants"]) == (207682, 207821, 14)
    assert e24["synthetic"] is False and e24["published"] is None
    assert e24["source_doc"] == "giridih_vs2024_form20.xlsx"
    assert e24["sources"][0]["booth_rows"] == 367 and e24["sources"][0]["sha256"]
    assert [(r["candidate"], r["booths"]) for r in e24["booths_led"]] == [
        ("Nirbhay Kumar Shahabadi", 193), ("Sudivya Kumar", 166), ("Navin Anand", 8)]
    assert rows["VS-2019"]["margin_votes"] == 15884
    assert body["data_health"]["form20_real_docs"] == 2
    assert body["data_health"]["form20_booth_rows"] == 734


def test_an_election_without_form20_keeps_its_published_result(api):
    rows = {e["label"]: e for e in _get(api, "/acs/32/summary")["elections"]}
    e14 = rows["VS-2014"]
    assert e14["has_results"] is False and e14["booths"] == 0
    assert e14["published"]["margin_votes"] == 9933
    assert e14["published"]["source"]


def test_every_candidate_of_2024_ranked_with_postal_and_deposit(api):
    body = _get(api, "/acs/32/elections/VS-2024/candidates")
    cands = body["candidates"]
    assert body["basis"] == "form20" and body["valid_votes"] == 207682
    assert len(cands) == 14 and [c["rank"] for c in cands] == list(range(1, 15))
    top = cands[0]
    assert (top["candidate"], top["party"], top["is_winner"]) == ("Sudivya Kumar", "JMM", True)
    assert top["evm_votes"] + top["postal_votes"] == top["votes"] == 94042
    assert top["booths_led"] == 166 and top["deposit_forfeited"] is False
    third = cands[2]
    assert (third["candidate"], third["party"], third["votes"]) == ("Navin Anand", "JLKM", 10787)
    assert third["deposit_forfeited"] is True        # 10,787 < 205,678 / 6
    assert sum(c["votes"] for c in cands) + body["nota"]["votes"] == 207682
    assert body["nota"]["votes"] == 2004
    unrecorded = [c for c in cands if not c["party_recorded"]]
    assert len(unrecorded) == 11 and all(c["party"] == "UNK" for c in unrecorded)
    assert sum(c["booths_led"] for c in cands) == 367
    assert "s.158" in body["deposit_rule"]
    assert body["sources"][0]["source_doc"] == "giridih_vs2024_form20.xlsx"


def test_candidates_of_an_unloaded_election_come_from_published_totals(api):
    body = _get(api, "/acs/32/elections/VS-2014/candidates")
    assert body["basis"] == "published"
    assert all(c["evm_votes"] is None and c["booths_led"] is None for c in body["candidates"])


def test_an_unknown_election_is_a_404(api):
    client, headers = api
    assert client.get("/acs/32/elections/VS-1999/candidates",
                      headers=headers["strategist"]).status_code == 404


def test_the_booth_card_names_every_candidate(api, real):
    uid = one(real["conn"], "SELECT booth_uid FROM booth WHERE current_ps_number = 1")["booth_uid"]
    card = _get(api, f"/acs/32/booths/{uid}/card")
    results = {r["election_label"]: r for r in card["results"]}
    assert len(results["VS-2024"]["candidates"]) == 15          # 14 and NOTA
    assert len(results["VS-2019"]["candidates"]) == 13
    assert results["VS-2024"]["tendered"] == 0
    names = [c["candidate"] for c in results["VS-2024"]["candidates"]]
    assert results["VS-2024"]["winner_candidate"] == names[0]


def test_the_booth_table_names_winner_and_runner_up(api):
    body = _get(api, "/acs/32/results/VS-2024/booths")
    assert body["count"] == 367
    row = body["rows"][0]
    assert row["winner_candidate"] and row["runner_candidate"]
    assert row["winner_candidate"] != row["runner_candidate"]
    assert 2 <= row["contestants"] <= 14 and row["tendered"] == 0     # who polled here


def test_the_block_user_sees_the_unassigned_block(api):
    body = _get(api, "/acs/32/results/VS-2024/booths", role="block")
    assert body["count"] == 367


def test_profiles_leave_out_nota(api):
    rows = _get(api, "/acs/32/candidates")["rows"]
    assert rows and all(r["party"] != "NOTA" for r in rows)

