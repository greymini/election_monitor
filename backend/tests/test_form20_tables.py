"""ingest/form20_tables.py on the real Giridih Form 20 workbooks (db/seed/form20/)."""

from __future__ import annotations

import shutil

import pytest

from ingest import form20_tables as f
from tests.paths import BACKEND

FORM20 = BACKEND / "db" / "seed" / "form20"
FILES = {2019: FORM20 / "giridih_vs2019_form20.xlsx", 2024: FORM20 / "giridih_vs2024_form20.xlsx"}

# The declared results, from each file's "Total Votes Polled" row (postal included).
PUBLISHED = {
    2024: {"SUDIVYA KUMAR": 94042, "NIRBHAY KUMAR SHAHABADI": 90204, "NAVIN ANAND": 10787,
           "SUNAINA PATHAK": 2635, "RAMESHWAR DUSADH": 2163, "AJEET RAY": 1540,
           "ARUNDHATI MISHRA": 1208, "ANISHA SINHA": 858, "DR. BARNABAS HEMBROM": 842,
           "ARTI DEVI": 306, "OM PRAKASH MAHTO": 301, "QAISAR JAMAL ANSARI": 289,
           "KRANTI KUMAR MURMU": 253, "ASHWINI AMBEDKAR": 250},
    2019: {"SUDIVYA KUMAR": 80871, "NIRBHAY KUMAR SHAHABADI": 64987, "CHUNNU KANT": 6903,
           "MOHAMMAD LALLU": 2493, "UPENDRA KUMAR SHARMA": 2026, "SHYAM PRASAD BARNWAL": 1533,
           "SIKANDAR ALI": 1272, "SHWETA KUMARI": 1182, "MD SALMAN SAHAB": 1124,
           "BUDHAN HEMABRAM": 991, "RAJESH KU SINHA": 834, "SIMA KUMARI": 833},
}


@pytest.fixture(scope="module", params=[2019, 2024])
def table(request):
    return request.param, f.read(FILES[request.param])


def test_the_file_is_internally_consistent(table):
    _, t = table
    assert f.problems(t) == []


def test_every_polling_station_is_present_once(table):
    _, t = table
    assert sorted(b.ps_number for b in t.booths) == list(range(1, 368))


def test_totals_match_the_declared_result(table):
    year, t = table
    assert t.published_votes() == PUBLISHED[year]


def test_ac_level_figures(table):
    year, t = table
    expected = {2024: (205678, 2004, 139, 207821, 1894 + 11),
                2019: (165049, 2773, 71, 167893, 269 + 2)}[year]
    valid, nota, rejected, polled, postal = expected
    assert (t.polled.valid, t.polled.nota, t.polled.rejected, t.polled.total) == (
        valid, nota, rejected, polled)
    assert t.postal.valid + t.postal.nota == postal


def test_the_2019_page_stamp_rows_are_recovered():
    """The last station on pages 1-11 is interleaved with the 'Page' stamp."""
    t = f.read(FILES[2019])
    recovered = {b.ps_number: b for b in t.booths if b.recovered}
    assert sorted(recovered) == [32, 64, 96, 128, 160, 192, 224, 256, 288, 320, 352]
    ps32 = recovered[32]
    assert ps32.votes[0] == 2                        # cell read "1 2": page 1, value 2
    assert sum(ps32.votes) == ps32.valid == 219


def test_header_names_are_cleaned():
    t = f.read(FILES[2024])
    assert t.candidates[1] == "NIRBHAY KUMAR SHAHABADI"
    assert f.display_name("DR. BARNABAS HEMBROM") == "Dr. Barnabas Hembrom"
    assert f.name_key("Sudivya  Kumar") == f.name_key("SUDIVYA\nKUMAR")


@pytest.fixture
def tampered(tmp_path):
    from openpyxl import load_workbook

    def make(edit):
        path = tmp_path / "tampered.xlsx"
        shutil.copy(FILES[2024], path)
        wb = load_workbook(path)
        edit(wb)
        wb.save(path)
        return f.read(path)
    return make


def test_a_changed_vote_cell_is_caught(tampered):
    def edit(wb):
        ws = wb["Page_1"]
        ws.cell(row=4, column=4).value = "448"      # PS 1, second candidate: 447 -> 448
    issues = f.problems(tampered(edit))
    assert any("PS 1:" in i for i in issues)
    assert any("Total EVM Votes" in i for i in issues)


def test_a_missing_polling_station_is_caught(tampered):
    def edit(wb):
        wb["Page_1"].delete_rows(5)                  # drop PS 2
    issues = f.problems(tampered(edit))
    assert any("missing from the sequence" in i for i in issues)
